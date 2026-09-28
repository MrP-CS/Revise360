"""Original 40-mark end-of-topic assessment and teacher mark scheme."""
from pathlib import Path
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml.ns import qn
from copy import deepcopy
import json
ROOT=Path(__file__).resolve().parents[2]
source=Document(ROOT/'worksheets/BL_L01_LogicGateLab_Worksheet.docx')
def setp(p,t):
 if p.runs:
  p.runs[0].text=t
  for r in p.runs[1:]:r.text=''
 else:p.add_run(t)
def panel(d,index,title,text,lines=0):
 e=deepcopy(source.tables[index]._tbl);d._element.body.insert(-1,e)
 from docx.table import Table
 c=Table(e,d._body).cell(0,0)
 ps=list(c.paragraphs)
 for x in list(c._tc):
  if x.tag!=qn('w:tcPr'):c._tc.remove(x)
 for src,value in [(ps[0],title),(ps[1],text)]:
  p=deepcopy(src._p)
  from docx.text.paragraph import Paragraph
  para=Paragraph(p,c);setp(para,value)
  if src==ps[1] and index!=0:
   for run in para.runs:run.font.size=Pt(10.5);run.font.italic=False;run.font.bold=False;run.font.color.rgb=RGBColor.from_string('101828')
  c._tc.append(p)
 if lines:
  t=deepcopy(source.tables[6].cell(0,0).tables[0]._tbl);row=t.find(qn('w:tr'))
  for _ in range(lines-1):t.append(deepcopy(row))
  c._tc.append(t)
 from docx.oxml import OxmlElement
 c._tc.append(OxmlElement('w:p'));d.add_paragraph().paragraph_format.space_after=Pt(2)
 return c

def border(t):
 from docx.oxml import OxmlElement
 b=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  v=OxmlElement('w:'+side);v.set(qn('w:val'),'single');v.set(qn('w:sz'),'5');v.set(qn('w:color'),'D9DEE7');b.append(v)
 t._tbl.tblPr.append(b)

def new(title,subtitle):
 d=Document(ROOT/'worksheets/BL_L01_LogicGateLab_Worksheet.docx')
 for e in list(d._element.body):
  if e.tag!=qn('w:sectPr'):d._element.body.remove(e)
 panel(d,0,title,subtitle)
 setp(d.sections[0].footer.paragraphs[0],'Revise360  •  Topic 2.1 Algorithms  •  Original practice assessment')
 d.core_properties.author='Revise360';return d
Q=[
('Computational thinking',6,'A school is designing an equipment-loan system.\na) Define abstraction. [2]\nb) Give one detail the system should retain and one it can ignore. [2]\nc) Explain how decomposition helps develop the system. [2]',
'1a Keep relevant details (1) and remove unnecessary detail for a purpose (1).\n1b Suitable relevant detail, e.g. borrower ID / item ID / due date (1). Suitable irrelevant detail, e.g. favourite film (1).\n1c Split into smaller sub-problems (1), making the parts easier to develop, understand or test (1).'),
('Structure diagram',4,'Draw a structure diagram for the equipment-loan system. It must include Members, Loans and Returns. Break Loans into two appropriate smaller tasks. [4]',
'Whole system above the three named main parts (1). Correct parent-child links (1). Check availability beneath Loans (1). Record loan beneath Loans (1). Accept equivalent appropriate loan sub-problems.'),
('Searching',6,'The list is [2, 4, 6, 8, 10, 12, 14]. Target = 12.\na) For a left-to-right linear search, list the values checked and give the comparison count. [2]\nb) For binary search, use zero-based indices and round midpoints down. State the initial midpoint value, the new index range and the next checked value. [3]\nc) State the data requirement for binary search. [1]',
'3a Checks 2,4,6,8,10,12 (1), six comparisons (1).\n3b First midpoint index 3 has value 8 (1). Keep indices 4 to 6 (1). Next midpoint index 5 has value 12, so found (1).\n3c Data must be sorted by the search key (1).'),
('Sorting',10,'a) Show the first full left-to-right bubble-sort pass on [6, 2, 5, 1]. Show each swap, then identify the value in its final position. [4]\nb) Merge sorted lists [2, 8] and [3, 7]. State the first two comparisons and the final list. [3]\nc) Show the whole list after each insertion when insertion-sorting [5, 2, 4, 1]. [3]',
'4a [2,6,5,1] (1), [2,5,6,1] (1), [2,5,1,6] (1). Value 6 is in its final position (1).\n4b Compare 2 and 3, take 2 (1). Compare 8 and 3, take 3 (1). Final [2,3,7,8] (1).\n4c [2,5,4,1] (1), [2,4,5,1] (1), [1,2,4,5] (1).'),
('Create a flowchart',5,'Draw a flowchart that inputs a pupil age. If age is at least 12, output Allowed. Otherwise output Too young. Include a start, end, appropriate symbols and labelled branches. [5]',
'Input age in an input/output symbol (1). Decision age >= 12 in a diamond (1). True and false branches correctly labelled (1). Correct output on each branch in input/output symbols (1). Connected start and end terminals with sensible flow lines (1).'),
('Debug an algorithm',3,'A programmer wants the average of a and b. The line is average = a + b / 2.\na) State whether this is a syntax or logic error. [1]\nb) Correct the expression and explain why your correction works. [2]',
'6a Logic error (1).\n6b average = (a + b) / 2 (1). Parentheses ensure addition happens before division (1).'),
('Complete a trace table',6,'Trace this OCR-style pseudocode.\n\ntotal = 0\nfor i = 1 to 4\n    total = total + i\nnext i\nprint(total)\n\nRecord the initial total, the total after each iteration and the final output. [6]',
'Initial total 0 (1). Totals after iterations: 1 (1), 3 (1), 6 (1), 10 (1). Final output 10 (1). Accept equivalent trace conventions that clearly distinguish variable state from output.')]
assert sum(q[1] for q in Q)==40
stem='AL_L17_AlgorithmsAssessment'
d=new('Topic 2.1 Algorithms: assessment','Lesson 17  •  40 marks  •  Suggested time: 45 minutes')
panel(d,2,'Instructions','Work independently with the experience and notes closed. Answer every question. Show the steps in searches and sorts. Use the marks to judge how much detail to include. Questions are original practice, not an official OCR paper.')
for i in [0,1]:panel(d,6+i,f'{i+1}  {Q[i][0]}  [{Q[i][1]}]',Q[i][2],8 if i==1 else 6)
d.add_page_break();panel(d,8,'3  Searching  [6]',Q[2][2],12)
d.add_page_break();panel(d,9,'4  Sorting  [10]',Q[3][2],20)
d.add_page_break();panel(d,10,'5  Create a flowchart  [5]',Q[4][2],17);panel(d,11,'6  Debug an algorithm  [3]',Q[5][2],5)
d.add_page_break();c=panel(d,6,'7  Complete a trace table  [6]',Q[6][2])
t=d.add_table(rows=7,cols=3);border(t)
for c,s in zip(t.rows[0].cells,['Stage / i','total','Output']):c.text=s
for c,s in zip([r.cells[0] for r in t.rows[1:]],['Initial','1','2','3','4','After loop']):c.text=s
for r in t.rows:
 r.height=Inches(.33)
d.add_paragraph('Total: ______ / 40')
d.add_page_break();panel(d,15,'Reflection after feedback','Complete this page when your marked assessment is returned.')
t=d.add_table(rows=8,cols=4);border(t)
for c,s in zip(t.rows[0].cells,['Question','Available','My mark','Next lesson to revisit']):c.text=s
for row,(q,lesson) in zip(t.rows[1:],zip(Q,['1–3','3','7–8','9–10','11','15','16'])):
 for c,s in zip(row.cells,[q[0],str(q[1]),'',lesson]):c.text=s
panel(d,14,'My next step','Choose one missed question. Write a corrected answer and explain the mistake you will avoid next time.',8)
d.save(ROOT/'worksheets'/f'{stem}_Worksheet.docx')
m=new('Topic 2.1 assessment: mark scheme','40 marks  •  For teacher use after the assessment')
panel(m,2,'Marking guidance','Award each distinct point once. Accept equivalent correct pseudocode and explanations. Credit follow-through where a single earlier error is carried consistently into a later trace, unless that later mark assesses the same error. No GCSE grade boundaries are assigned to this topic test.')
for i,q in enumerate(Q):
 if i in [2,4]:m.add_page_break()
 panel(m,6+i%6,f'{i+1}  {q[0]}  [{q[1]}]',q[3])
m.save(ROOT/'worksheets'/f'{stem}_MarkScheme.docx')
Path(__file__).with_name('assessment.json').write_text(json.dumps(Q,indent=2))
print('Assessment and mark scheme written')
