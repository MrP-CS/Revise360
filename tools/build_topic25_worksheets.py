"""Original editable worksheets for the Topic 2.5 experiences. Requires python-docx."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from topic25_content import LESSONS
ROOT=Path(__file__).resolve().parents[1]
def paragraph(d,text='',bold=False,size=11):
 p=d.add_paragraph();r=p.add_run(text);r.bold=bold;r.font.size=Pt(size);p.paragraph_format.space_after=Pt(6);return p

def lines(d,n):
 for _ in range(n):
  p=d.add_paragraph('________________________________________________________________________')
  p.paragraph_format.space_after=Pt(7);p.runs[0].font.size=Pt(10);p.runs[0].font.color.rgb=RGBColor.from_string('BAC3CF')
def base(title,lesson):
 d=Document();s=d.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69);s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.7)

 for st in d.styles:
  for border in st.element.findall('.//' + qn('w:pBdr')): border.getparent().remove(border)
 normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.12
 for st in ['Title','Heading 1','Heading 2']:
  d.styles[st].font.name='Calibri';d.styles[st].font.color.rgb=RGBColor(0,0,0)
 d.styles['Title'].font.size=Pt(24);d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(13)
 h=s.header.paragraphs[0];h.text='Revise360  |  GCSE Computer Science  |  Topic 2.5';h.runs[0].font.size=Pt(9)
 f=s.footer.paragraphs[0];f.text=f'Original Revise360 resource  |  Lesson {lesson}  |  ';fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
 d.core_properties.author='Revise360';d.core_properties.title=title
 d.add_paragraph(title,'Title');paragraph(d,f'Topic 2.5  |  Lesson {lesson}',True)
 paragraph(d,'Name ____________________    Class __________    Date __________')
 return d

def station(d,k,s):
 d.add_paragraph(str(k+1)+' '+s['name'],'Heading 2');paragraph(d,'Key facts in your own words',True);lines(d,3);paragraph(d,s['challenge']);lines(d,2)
for l in LESSONS:
 d=base(l['title'],l['lesson']);paragraph(d,l['subtitle'])
 d.add_paragraph('Learning outcomes','Heading 1')
 for o in l['objectives']:paragraph(d,o)
 d.add_paragraph('Starter','Heading 1');paragraph(d,l['starter']);lines(d,3)
 d.add_paragraph('Key terms','Heading 1')
 for term,_ in l['terms']:paragraph(d,term,True);lines(d,2)
 d.add_paragraph('Open the experience','Heading 1');paragraph(d,'Open '+l['title']+' from Topic 2.5. Read the wall panels and blue fact markers. Write the key facts in your own words, explain your answers and complete the numbered activities. Use the star for the plenary knowledge check toward the end of the lesson.')
 if l['lesson']==2:paragraph(d,'The invented instruction set is enrichment. You do not need to memorise it for OCR J277.',True)
 d.add_page_break();d.add_paragraph('Explore and explain','Title');paragraph(d,'Stations 1 and 2: turn right. Station 3: turn around. Record key facts in your own words.')
 for k in range(3):station(d,k,l['stations'][k])
 d.add_page_break();d.add_paragraph('Apply and investigate','Title');paragraph(d,'Station 4: turn around. Stations 5 and 6: turn left. Use the interactive board where provided.')
 for k in range(3,6):station(d,k,l['stations'][k])
 d.add_page_break();d.add_paragraph('Plenary knowledge check','Title');paragraph(d,'Complete the star station, then close the experience and answer these questions from memory. Use accurate technical terms and link explanations to the situation.')
 for k,(question,marks,answer) in enumerate(l['exam']):
  paragraph(d,f'{k+1}. {question} [{marks} marks]',True);lines(d,6 if marks>=4 else 4)
 d.add_paragraph('Review your learning','Heading 1');paragraph(d,'Experience score ______ / ______     Written plenary ______ / '+str(sum(q[1] for q in l['exam'])))
 paragraph(d,'One correction I can now explain');lines(d,2);paragraph(d,'My next step or question');lines(d,2)
 stem=f'PL_L{l["lesson"]:02d}_{l["slug"]}';d.save(ROOT/f'worksheets/{stem}_Worksheet.docx')
assessment=[
 ('A conservation charity needs an app that several programmers can update and adapt for different computers. Explain two reasons to choose a high-level language.',4,'Readable expressions/meaningful source (1) make maintenance easier for a team (1). Source portability (1) helps target different systems with suitable tools (1). Accept other relevant explained reasons.'),
 ('Explain why machine code for one type of processor may not execute on a different type.',2,'Machine code follows a particular instruction set (1); the other processor may not recognise the instructions or their encoding (1).'),
 ('Explain why high-level source needs a translator.',2,'The processor executes its own machine instructions (1); a translator processes high-level source into executable code or operations (1).'),
 ('Describe two differences between a compiler and an interpreter in the GCSE model.',4,'Whole-source translation before execution versus processing during execution (2 for the paired contrast). Saved target code can be reused versus source processed on each run (2). Accept valid differences about build/error behaviour with the appropriate model.'),
 ('A developer wants to distribute a stable program to users on a specified platform without giving them the high-level source. Recommend an approach and explain your choice.',2,'Compiler (1); it creates a target-code build that can be distributed without requiring users to translate the high-level source (1).'),
 ('Describe how an editor and error diagnostics each help a programmer develop a program.',4,'Editor creates/changes source (1), allowing instructions or corrections to be entered (1). Diagnostics identify/explain a problem or location (1), helping investigation and correction (1).'),
 ('An order has price 6 and quantity 4. The program uses total = price + quantity and outputs 10. Explain why a syntax check might not flag this and state the correction.',2,'Addition is syntactically valid but implements the wrong calculation/logic (1); use multiplication, giving 24 (1).')]
d=base('Programming languages and IDEs assessment',5);paragraph(d,'20 marks  |  Suggested time 25 minutes  |  Work independently')
paragraph(d,'These are original exam-style questions, not an OCR examination paper. Write clear explanations and use the mark values to judge the detail needed.')
for k,(q,m,a) in enumerate(assessment):
 if k==3:d.add_page_break();d.add_paragraph('Assessment continued','Title')
 paragraph(d,f'{k+1}. {q} [{m} marks]',True);lines(d,5 if m==4 else 3)
paragraph(d,'Total ______ / 20     A topic I will revisit __________________________',True)
d.save(ROOT/'worksheets/PL_L05_EndOfTopic_Assessment.docx')
p=ROOT/'docs/topic25-trial-guide.md';s=p.read_text().split('\n## Original paper assessment')[0];s+='\n## Original paper assessment marking guide\n\nTotal 20 marks. Award equivalent accurate explanations; do not require memorised wording.\n'
for k,(q,m,a) in enumerate(assessment):s+=f'\n### Question {k+1} [{m}]\n\n{q}\n\n{a}\n'
p.write_text(s)
print('Created four lesson worksheets and one original 20-mark assessment.')
