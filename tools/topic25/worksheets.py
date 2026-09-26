from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[2];ls=json.loads((Path(__file__).parent/'lessons.json').read_text())
def p(d,s='',style=None):return d.add_paragraph(s,style)
def lines(d,n=2):
 for _ in range(n):
  x=p(d,'________________________________________________________________________________');x.paragraph_format.space_after=Pt(5);x.runs[0].font.color.rgb=RGBColor.from_string('A7B1C0');x.runs[0].font.size=Pt(10)
def heading(d,s):d.add_heading(s,2)
def page(d,l,title):
 d.add_page_break();d.add_heading(title,1);p(d,f'Topic 2.5  •  Lesson {l["n"]}  •  {l["title"]}')
def task(d,t,n):
 p(d,f'{n}. {t["q"]}').runs[0].bold=True
 if t['t']=='sort':
  p(d,'Categories: '+ ' / '.join(t['cats']))
  for name,_ in t['items']:p(d,name+'  ____________________')
 elif t['t']=='match':
  p(d,'Choices: '+' / '.join(y for _,y in reversed(t['pairs'])))
  for x,_ in t['pairs']:p(d,x+'  ____________________')
 elif t['t']=='multi':
  for v in t['opts']:p(d,'[  ] '+v)
 else:
  for v in t['a'][1:]+t['a'][:1]:p(d,'[  ] '+v)
 lines(d,1)
for l in ls:
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.65)
 st=d.styles['Normal'];st.font.name='DejaVu Sans';st.font.size=Pt(10);st.paragraph_format.space_after=Pt(6)
 for name,size in [('Title',25),('Heading 1',20),('Heading 2',13)]:
  st=d.styles[name];st.font.name='DejaVu Sans';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('000000'if name=='Title'else'2461CF')
 foot=sec.footer.paragraphs[0];foot.text='Revise360  •  Topic 2.5  •  ';fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld);foot.runs[0].font.size=Pt(9)
 d.add_heading(l['title'],0);p(d,f'Topic 2.5 Programming languages and IDEs  |  Lesson {l["n"]}');p(d,'Name: ____________________    Class: __________    Date: __________')
 heading(d,'Do Now research and discuss');p(d,l['donow']);lines(d,2)
 heading(d,'Learning outcomes')
 for x in l['outcomes']:p(d,'• '+x)
 heading(d,'Starter');p(d,l['starter']);lines(d,3)
 heading(d,'Key terminology quiz');p(d,'Write a definition for each term.')
 for term,_ in l['terms']:p(d,term).runs[0].bold=True;lines(d,2)
 for group in range(2):
  page(d,l,f'Experience stations {group*3+1} to {group*3+3}')
  if group==0:p(d,f'Open {l["title"]} in Topic 2.5. Read each panel or blue information marker. Write the key facts IN YOUR OWN WORDS and answer the challenge before tapping the numbered badge.')
  for i in range(group*3,group*3+3):
   s=l['stations'][i];heading(d,f'{i+1} {s["name"]}');p(d,'Key facts in my own words');lines(d,2);p(d,s['challenge']);lines(d,2)
 page(d,l,'Final challenge and reflection');p(d,'Attempt the tasks below before selecting the star in the experience.')
 for i,t in enumerate(l['final']):task(d,t,i+1)
 p(d,'Experience score: ______ / ______');p(d,'Use review mode to correct mistakes. Explain one correction below. Your first-attempt score stays the same.');lines(d,2)
 p(d,'Return to the Do Now question. What can you now add to your answer?');lines(d,2)
 page(d,l,'Plenary knowledge checks');p(d,'Near the end of the lesson, close the experience and answer from memory. Use the marks to guide the detail you include.')
 for i,(question,answer) in enumerate(l['plenary']):
  heading(d,chr(97+i));p(d,question);marks=int(re.search(r'\[(\d+)\]',question)[1]);lines(d,marks+1)
 heading(d,'Confidence check')
 for x in l['outcomes']:
  p(d,x);p(d,'[ ] Not yet    [ ] Getting there    [ ] Confident')
 p(d,'One thing I will revisit: ___________________________________________________')
 d.core_properties.author='Revise360';d.core_properties.title=l['title'];d.core_properties.subject='Original Topic 2.5 lesson worksheet'
 for root in [d._element,d.styles.element]:
  for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
 d.save(ROOT/'worksheets'/f'{l["stem"]}_Worksheet.docx')
 print(l['stem'])
